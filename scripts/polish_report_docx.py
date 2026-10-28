"""
Post-processing pass for Bao_Cao_Do_An_UFM_Rec_Nhom_4.docx.

Two jobs:
1. Replace every inline LaTeX span ("$...$") and every leftover LaTeX command
   with plain Unicode so the Word file shows no raw markup.
2. Replace the LaTeX display equations inside the callout boxes with native
   Word equations (OMML), taken from scripts/report_equations_omml.json.
"""

import json
import re
from pathlib import Path

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(__file__).resolve().parents[1]
EQUATION_RESOURCE = BASE_DIR / "scripts" / "report_equations_omml.json"

GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
    "zeta": "ζ", "eta": "η", "theta": "θ", "iota": "ι", "kappa": "κ",
    "lambda": "λ", "mu": "μ", "nu": "ν", "xi": "ξ", "pi": "π", "rho": "ρ",
    "sigma": "σ", "tau": "τ", "upsilon": "υ", "phi": "φ", "chi": "χ",
    "psi": "ψ", "omega": "ω", "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ",
    "Lambda": "Λ", "Xi": "Ξ", "Pi": "Π", "Sigma": "Σ", "Phi": "Φ",
    "Psi": "Ψ", "Omega": "Ω",
}

SYMBOLS = {
    "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "approx": "≈",
    "neq": "≠", "ne": "≠", "equiv": "≡", "sim": "∼", "simeq": "≃",
    "in": "∈", "notin": "∉", "subset": "⊂", "subseteq": "⊆",
    "to": "→", "rightarrow": "→", "leftarrow": "←", "Rightarrow": "⇒",
    "Leftarrow": "⇐", "implies": "⇒", "iff": "⇔", "mapsto": "↦",
    "cdot": "·", "times": "×", "div": "÷", "pm": "±", "mp": "∓",
    "ll": "≪", "gg": "≫", "propto": "∝", "partial": "∂", "nabla": "∇",
    "infty": "∞", "dots": "…", "cdots": "⋯", "ldots": "…",
    "sum": "∑", "prod": "∏", "int": "∫", "sqrt": "√", "oplus": "⊕",
    "otimes": "⊗", "odot": "⊙", "circ": "∘", "cup": "∪", "cap": "∩", "forall": "∀",
    "exists": "∃", "perp": "⊥", "parallel": "∥", "prime": "′",
    "ldots": "…", "quad": " ", "qquad": "  ",
}

BLACKBOARD = {"R": "ℝ", "N": "ℕ", "Z": "ℤ", "Q": "ℚ", "C": "ℂ"}

ACCENTS = {"hat": "̂", "bar": "̄", "tilde": "̃", "vec": "⃗"}


def _replace_accent(match):
    accent = ACCENTS[match.group(1)]
    return match.group(2) + accent


def _replace_blackboard(match):
    return BLACKBOARD.get(match.group(1), match.group(1))


def _replace_text_cmd(match):
    return match.group(1)


def _replace_symbol(match):
    name = match.group(1)
    if name in GREEK:
        return GREEK[name]
    return SYMBOLS.get(name, match.group(0))


def latex_fragment_to_unicode(fragment, strip_grouping=False):
    fragment = re.sub(r"\\(?:text|mathrm|mathit|mathbf)\{([^{}]*)\}", r"\1", fragment)
    fragment = re.sub(r"\\(hat|bar|tilde|vec)\{([^{}]*)\}", _replace_accent, fragment)
    fragment = re.sub(r"\\mathbb\{([^{}]*)\}", _replace_blackboard, fragment)
    fragment = re.sub(r"\\mathcal\{([^{}]*)\}", r"\1", fragment)
    fragment = re.sub(r"\\([A-Za-z]+)", _replace_symbol, fragment)
    if strip_grouping:
        # LaTeX grouping braces are never meant to be displayed; escaped ones are.
        fragment = re.sub(r"(?<!\\)[{}]", "", fragment)
    fragment = fragment.replace(r"\{", "{").replace(r"\}", "}")
    fragment = fragment.replace(r"\|", "‖").replace(r"\_", "_")
    fragment = fragment.replace(r"\%", "%").replace(r"\$", "$").replace(r"\&", "&")
    fragment = fragment.replace("\\\\", "; ")
    fragment = fragment.replace("^*", "*")
    fragment = re.sub(r"\\[^A-Za-z]", "", fragment)
    return fragment


def latex_to_unicode(text):
    """Convert inline $...$ spans, then any bare LaTeX commands left behind."""
    text = re.sub(
        r"\$([^$]+)\$",
        lambda m: latex_fragment_to_unicode(m.group(1), strip_grouping=True),
        text,
    )
    if "\\" in text or re.search(r"[\^_]\{", text):
        text = latex_fragment_to_unicode(text, strip_grouping=True)
    return text


def _iter_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _replace_run_text(paragraph, transform):
    for run in paragraph.runs:
        replaced = transform(run.text)
        if replaced != run.text:
            run.text = replaced


def _replace_display_equation(doc, equations):
    """Swap the LaTeX equation paragraph of each callout box for native OMML."""
    replaced = []
    for table in doc.tables:
        if len(table.rows) != 1 or len(table.columns) != 1:
            continue
        cell = table.cell(0, 0)
        paragraphs = cell.paragraphs
        if len(paragraphs) < 3:
            continue
        match = re.match(r"Công thức\s+([\d.]+)\.", paragraphs[0].text.strip())
        if not match:
            continue
        formula_id = match.group(1)
        omml = equations.get(formula_id)
        if not omml:
            continue
        old_p = paragraphs[1]._p
        new_p = parse_xml(
            f'<w:p {nsdecls("w", "m")}>'
            f'<w:pPr><w:jc w:val="center"/></w:pPr>'
            f"{omml}</w:p>"
        )
        old_p.addprevious(new_p)
        old_p.getparent().remove(old_p)
        replaced.append(formula_id)
    return replaced


def polish(doc_path, output_path=None):
    doc = docx.Document(str(doc_path))
    equations = json.loads(EQUATION_RESOURCE.read_text(encoding="utf-8"))

    for paragraph in _iter_paragraphs(doc):
        _replace_run_text(paragraph, latex_to_unicode)

    replaced = _replace_display_equation(doc, equations)

    output_path = output_path or doc_path
    doc.save(str(output_path))
    return replaced


if __name__ == "__main__":
    import sys

    target = Path(sys.argv[1]) if len(sys.argv) > 1 else BASE_DIR / "Bao_Cao_Do_An_UFM_Rec_Nhom_4.docx"
    ids = polish(target)
    print(f"Polished {target.name}: {len(ids)} native equations -> {', '.join(ids)}")
